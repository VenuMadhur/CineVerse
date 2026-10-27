"""
Tests for the Operational Reports API (Module 3, Phase 3d):
retry queue, retry success, retry failure, cancelled due to payment
failure.
"""
import csv
import io
import uuid
from unittest.mock import patch


def _detail_rows(csv_text, header):
    """Mirrors the exact CSV-parsing pattern established in
    test_tax_reports_api.py (Phase 2b)."""
    all_rows = list(csv.reader(io.StringIO(csv_text)))
    assert header in all_rows, f"header {header} not found in CSV"
    start = all_rows.index(header) + 1
    return all_rows[start:]


_COLUMNS = ["Customer", "Invoice", "Retry Attempt", "Status", "Scheduled Date", "Failure Reason"]


def test_operational_report_requires_admin(client, auth_headers):
    resp = client.get(
        "/api/v1/reports/operational", headers=auth_headers, params={"dimension": "retry_queue"}
    )
    assert resp.status_code == 403


def test_operational_report_invalid_dimension_rejected(client, admin_auth_headers):
    resp = client.get(
        "/api/v1/reports/operational", headers=admin_auth_headers, params={"dimension": "bogus"}
    )
    assert resp.status_code == 422


def _configure_retry(client, admin_auth_headers):
    client.post(
        "/api/v1/retry/configuration",
        headers=admin_auth_headers,
        json={"retry_attempt": 1, "retry_after_days": 1},
    )


def test_retry_queue_shows_entry_after_failure(
    client, auth_headers, admin_auth_headers, active_subscription_with_invoice
):
    _configure_retry(client, admin_auth_headers)
    fixture = active_subscription_with_invoice
    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        client.post(
            "/api/v1/payments/process", headers=auth_headers, json={"invoice_id": fixture["invoice_id"]}
        )

    resp = client.get(
        "/api/v1/reports/operational",
        headers=admin_auth_headers,
        params={"dimension": "retry_queue", "format": "csv"},
    )
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), _COLUMNS)
    assert any(r[0] == fixture["customer"]["name"] for r in rows)
    assert any(r[3] == "pending" for r in rows)


def test_retry_failure_report_shows_failed_attempt(
    client, auth_headers, admin_auth_headers, active_subscription_with_invoice, db_session
):
    _configure_retry(client, admin_auth_headers)
    fixture = active_subscription_with_invoice
    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        client.post(
            "/api/v1/payments/process", headers=auth_headers, json={"invoice_id": fixture["invoice_id"]}
        )

    # Backdate the queue entry and run the scheduler for real, so it
    # actually attempts (and, forced to fail again) transitions to FAILED.
    from datetime import datetime, timedelta, timezone

    from app.models.retry_queue import RetryQueue
    from app.scheduler.retry_scheduler_service import RetrySchedulerService

    entry = (
        db_session.query(RetryQueue)
        .filter(RetryQueue.invoice_id == uuid.UUID(fixture["invoice_id"]))
        .first()
    )
    entry.scheduled_retry_date = datetime.now(timezone.utc) - timedelta(minutes=1)
    db_session.commit()

    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        RetrySchedulerService(db_session).run()

    resp = client.get(
        "/api/v1/reports/operational",
        headers=admin_auth_headers,
        params={"dimension": "retry_failure", "format": "csv"},
    )
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), _COLUMNS)
    assert any(r[3] == "failed" for r in rows)


def test_retry_success_report_shows_succeeded_attempt(
    client, auth_headers, admin_auth_headers, active_subscription_with_invoice, db_session
):
    _configure_retry(client, admin_auth_headers)
    fixture = active_subscription_with_invoice
    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        client.post(
            "/api/v1/payments/process", headers=auth_headers, json={"invoice_id": fixture["invoice_id"]}
        )

    from datetime import datetime, timedelta, timezone

    from app.models.retry_queue import RetryQueue
    from app.scheduler.retry_scheduler_service import RetrySchedulerService

    entry = (
        db_session.query(RetryQueue)
        .filter(RetryQueue.invoice_id == uuid.UUID(fixture["invoice_id"]))
        .first()
    )
    entry.scheduled_retry_date = datetime.now(timezone.utc) - timedelta(minutes=1)
    db_session.commit()

    with patch("app.services.payment_gateway_service.random.random", return_value=0.0):
        RetrySchedulerService(db_session).run()

    resp = client.get(
        "/api/v1/reports/operational",
        headers=admin_auth_headers,
        params={"dimension": "retry_success", "format": "csv"},
    )
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), _COLUMNS)
    assert any(r[3] == "succeeded" for r in rows)


def test_cancelled_due_to_failure_report(
    client, auth_headers, admin_auth_headers, active_subscription_with_invoice, db_session
):
    """Only one configured attempt -- the first failed retry exhausts
    the schedule and cancels the subscription."""
    client.post(
        "/api/v1/retry/configuration",
        headers=admin_auth_headers,
        json={"retry_attempt": 1, "retry_after_days": 1},
    )
    fixture = active_subscription_with_invoice
    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        client.post(
            "/api/v1/payments/process", headers=auth_headers, json={"invoice_id": fixture["invoice_id"]}
        )

    from datetime import datetime, timedelta, timezone

    from app.models.retry_queue import RetryQueue
    from app.scheduler.retry_scheduler_service import RetrySchedulerService

    entry = (
        db_session.query(RetryQueue)
        .filter(RetryQueue.invoice_id == uuid.UUID(fixture["invoice_id"]))
        .first()
    )
    entry.scheduled_retry_date = datetime.now(timezone.utc) - timedelta(minutes=1)
    db_session.commit()

    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        RetrySchedulerService(db_session).run()

    resp = client.get(
        "/api/v1/reports/operational",
        headers=admin_auth_headers,
        params={"dimension": "cancelled_due_to_failure", "format": "csv"},
    )
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), ["Customer", "Email", "Plan", "Cancelled Date"])
    assert any(r[0] == fixture["customer"]["name"] for r in rows)


def test_operational_report_xlsx_and_pdf_formats(client, admin_auth_headers):
    from openpyxl import load_workbook

    xlsx_resp = client.get(
        "/api/v1/reports/operational",
        headers=admin_auth_headers,
        params={"dimension": "retry_queue", "format": "xlsx"},
    )
    assert xlsx_resp.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_resp.content))
    assert "Detail" in workbook.sheetnames

    pdf_resp = client.get(
        "/api/v1/reports/operational",
        headers=admin_auth_headers,
        params={"dimension": "retry_queue", "format": "pdf"},
    )
    assert pdf_resp.status_code == 200
    assert pdf_resp.content[:4] == b"%PDF"

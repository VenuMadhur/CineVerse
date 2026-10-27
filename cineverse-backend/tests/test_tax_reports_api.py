"""
Tests for the Tax Reports API (Module 3, Phase 2b): all 9 report
views (daily/weekly/monthly/yearly collection, country, state, plan,
customer, payment method), export formats, and admin-gating.

CSV parsing follows the exact same pattern as test_reports_api.py:
the export is a multi-section document (title, summary block, blank
line, THEN the detail table's own header + rows) -- not a single flat
table -- so tests search for the column header row within the full
parsed CSV rather than treating row 1 as the header (which is what
report.title actually is).
"""
import csv
import io
import uuid
from unittest.mock import patch

from openpyxl import load_workbook

from app.scheduler.billing_scheduler_service import BillingSchedulerService

ALL_DIMENSIONS = [
    "daily",
    "weekly",
    "monthly",
    "yearly",
    "country",
    "state",
    "plan",
    "customer",
    "payment_method",
]


def _seed_paid_invoice_with_tax(client, admin_auth_headers, db_session, country="US", state="California"):
    """A real customer with a real paid invoice carrying a non-zero
    tax_amount (US is in the seeded Tax Master table at 8%, or falls
    back to settings.TAX_RATES identically if unseeded)."""
    email = f"taxreport_{uuid.uuid4().hex[:8]}@test.com"
    client.post(
        "/api/v1/auth/register",
        json={"email": email, "full_name": "Tax Report Data", "password": "SuperSecret123", "country": country},
    )
    token = client.post(
        "/api/v1/auth/login", json={"email": email, "password": "SuperSecret123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    plan = client.post(
        "/api/v1/plans",
        headers=admin_auth_headers,
        json={
            "name": f"Tax Report Plan {uuid.uuid4().hex[:6]}",
            "price": 100,
            "billing_interval": "monthly",
            "trial_period_days": 0,
        },
    ).json()
    customer = client.get("/api/v1/customers/me", headers=headers).json()
    if state:
        customer = client.patch(
            f"/api/v1/customers/{customer['id']}",
            headers=headers,
            json={"state": state}
        ).json()
    sub = client.post(
        "/api/v1/subscriptions",
        headers=headers,
        json={"customer_id": customer["id"], "plan_id": plan["id"]},
    ).json()
    client.post(f"/api/v1/subscriptions/{sub['id']}/activate", headers=headers)
    # Issue 2 fix: trial_period_days=0 now starts the subscription
    # ACTIVE directly, so /activate above is a no-op; the scheduler is
    # what invoices it (same fix as the other analytics/reports test
    # files touched in this pass).
    BillingSchedulerService(db_session).run()
    invoice_id = client.get(
        "/api/v1/invoices", params={"customer_id": customer["id"]}, headers=headers
    ).json()["items"][0]["id"]

    with patch("app.services.payment_gateway_service.random.random", return_value=0.0):
        client.post(
            "/api/v1/payments/process", headers=headers, json={"invoice_id": invoice_id}
        )
    return customer, plan, invoice_id


def _detail_rows(csv_text, header):
    """Returns the data rows that follow `header` (a list of column
    names) somewhere in the full CSV, as a list of lists -- mirrors
    the exact parsing pattern test_reports_api.py already uses."""
    all_rows = list(csv.reader(io.StringIO(csv_text)))
    assert header in all_rows, f"header {header} not found in CSV"
    start = all_rows.index(header) + 1
    return all_rows[start:]


def test_tax_report_requires_admin(client, auth_headers):
    resp = client.get(
        "/api/v1/reports/tax", headers=auth_headers, params={"dimension": "country"}
    )
    assert resp.status_code == 403


def test_tax_report_requires_auth(client):
    resp = client.get("/api/v1/reports/tax", params={"dimension": "country"})
    assert resp.status_code == 401


def test_tax_report_invalid_dimension_rejected(client, admin_auth_headers):
    resp = client.get(
        "/api/v1/reports/tax", headers=admin_auth_headers, params={"dimension": "bogus"}
    )
    assert resp.status_code == 422


def test_all_nine_dimensions_return_csv(client, admin_auth_headers, db_session):
    _seed_paid_invoice_with_tax(client, admin_auth_headers, db_session)
    for dimension in ALL_DIMENSIONS:
        resp = client.get(
            "/api/v1/reports/tax",
            headers=admin_auth_headers,
            params={"dimension": dimension, "format": "csv"},
        )
        assert resp.status_code == 200, dimension
        assert resp.headers["content-type"].startswith("text/csv")
        text = resp.content.decode()
        assert "Tax Report" in text
        assert "Total Tax Collected" in text


def test_tax_report_xlsx_and_pdf_formats(client, admin_auth_headers, db_session):
    _seed_paid_invoice_with_tax(client, admin_auth_headers, db_session)

    xlsx_resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={"dimension": "country", "format": "xlsx"},
    )
    assert xlsx_resp.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_resp.content))
    assert "Detail" in workbook.sheetnames
    assert workbook["Detail"].max_row >= 2  # header + at least 1 data row

    pdf_resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={"dimension": "country", "format": "pdf"},
    )
    assert pdf_resp.status_code == 200
    assert pdf_resp.content[:4] == b"%PDF"


def test_tax_by_country_reflects_seeded_rate(client, admin_auth_headers, db_session):
    _seed_paid_invoice_with_tax(client, admin_auth_headers, db_session, country="US")
    resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={"dimension": "country", "format": "csv"},
    )
    rows = _detail_rows(resp.content.decode(), ["Country", "Tax Collected"])
    us_row = next((r for r in rows if r[0] == "US"), None)
    assert us_row is not None
    # 8% of a $100 plan = $8.00 tax, whether sourced from the Tax
    # Master table or the settings fallback (both agree, Phase 2a).
    assert float(us_row[1]) >= 8.00


def test_tax_by_state_only_includes_customers_with_state_set(client, admin_auth_headers, db_session):
    _seed_paid_invoice_with_tax(client, admin_auth_headers, db_session, country="US", state="Texas")
    resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={"dimension": "state", "format": "csv"},
    )
    rows = _detail_rows(resp.content.decode(), ["State", "Tax Collected"])
    assert any(r[0] == "Texas" for r in rows)


def test_tax_by_plan_and_customer_and_payment_method(client, admin_auth_headers, db_session):
    customer, plan, _invoice_id = _seed_paid_invoice_with_tax(client, admin_auth_headers, db_session)

    plan_resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={"dimension": "plan", "format": "csv"},
    )
    plan_rows = _detail_rows(plan_resp.content.decode(), ["Plan", "Tax Collected"])
    assert any(r[0] == plan["name"] for r in plan_rows)

    customer_resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={"dimension": "customer", "format": "csv"},
    )
    customer_rows = _detail_rows(customer_resp.content.decode(), ["Customer", "Tax Collected"])
    assert any(r[0] == customer["name"] for r in customer_rows)

    method_resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={"dimension": "payment_method", "format": "csv"},
    )
    method_rows = _detail_rows(
        method_resp.content.decode(), ["Payment Method", "Tax Collected"]
    )
    assert len(method_rows) >= 1


def test_tax_report_date_range_filtering(client, admin_auth_headers, db_session):
    _seed_paid_invoice_with_tax(client, admin_auth_headers, db_session)
    resp = client.get(
        "/api/v1/reports/tax",
        headers=admin_auth_headers,
        params={
            "dimension": "monthly",
            "format": "csv",
            "date_from": "2020-01-01",
            "date_to": "2020-01-31",
        },
    )
    # A date range entirely before the seeded invoice was paid should
    # correctly exclude it -- proving date_from/date_to are genuinely
    # applied, not ignored.
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), ["Period", "Tax Collected"])
    assert rows == []

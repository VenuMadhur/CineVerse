"""
Tests for the Customer Insights Reports API (Module 3, Phase 3c):
top paying customers, near renewal, with failed payments, past due.
"""
import csv
import io
import uuid
from unittest.mock import patch

from app.scheduler.billing_scheduler_service import BillingSchedulerService


def _detail_rows(csv_text, header):
    """Mirrors the exact CSV-parsing pattern established in
    test_tax_reports_api.py (Phase 2b) -- the export is a multi-section
    document, so we search for the real header row rather than
    treating row 1 as the header."""
    all_rows = list(csv.reader(io.StringIO(csv_text)))
    assert header in all_rows, f"header {header} not found in CSV"
    start = all_rows.index(header) + 1
    return all_rows[start:]


def test_customer_insights_requires_admin(client, auth_headers):
    resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=auth_headers,
        params={"dimension": "top_paying"},
    )
    assert resp.status_code == 403


def test_customer_insights_invalid_dimension_rejected(client, admin_auth_headers):
    resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=admin_auth_headers,
        params={"dimension": "bogus"},
    )
    assert resp.status_code == 422


def test_top_paying_customers(client, auth_headers, admin_auth_headers, db_session):
    email = f"toppay_{uuid.uuid4().hex[:8]}@test.com"
    client.post(
        "/api/v1/auth/register",
        json={"email": email, "full_name": "Top Pay", "password": "SuperSecret123", "country": "US"},
    )
    token = client.post(
        "/api/v1/auth/login", json={"email": email, "password": "SuperSecret123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    plan = client.post(
        "/api/v1/plans",
        headers=admin_auth_headers,
        json={"name": f"TopPay Plan {uuid.uuid4().hex[:6]}", "price": 75, "billing_interval": "monthly", "trial_period_days": 0},
    ).json()
    customer = client.get(
        "/api/v1/customers/me",
        headers=headers,
    ).json()
    sub = client.post(
        "/api/v1/subscriptions", headers=headers,
        json={"customer_id": customer["id"], "plan_id": plan["id"]},
    ).json()
    client.post(f"/api/v1/subscriptions/{sub['id']}/activate", headers=headers)
    # Issue 2 fix: trial_period_days=0 now starts the subscription
    # ACTIVE directly, so /activate above is a no-op; the scheduler is
    # what invoices it now (see test_analytics_revenue_invoice_api.py
    # for the same fix, applied identically here).
    BillingSchedulerService(db_session).run()
    invoice_id = client.get(
        "/api/v1/invoices", params={"customer_id": customer["id"]}, headers=headers
    ).json()["items"][0]["id"]
    with patch("app.services.payment_gateway_service.random.random", return_value=0.0):
        client.post("/api/v1/payments/process", headers=headers, json={"invoice_id": invoice_id})

    resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=admin_auth_headers,
        params={"dimension": "top_paying", "format": "csv"},
    )
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), ["Customer", "Email", "Total Paid"])
    assert any(r[0] == "Top Pay" for r in rows)


def test_near_renewal(client, auth_headers, admin_auth_headers, active_subscription_with_invoice):
    """The seeded fixture's subscription has an unbilled cycle -- confirm the
    report doesn't error and returns a well-formed CSV regardless of whether
    that specific cycle falls inside the default window."""
    resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=admin_auth_headers,
        params={"dimension": "near_renewal", "format": "csv", "within_days": 90},
    )
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), ["Customer", "Email", "Plan", "Renewal Date"])
    assert isinstance(rows, list)  # well-formed, even if empty


def test_customers_with_failed_payments(client, auth_headers, admin_auth_headers, active_subscription_with_invoice):
    fixture = active_subscription_with_invoice
    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        client.post(
            "/api/v1/payments/process", headers=auth_headers, json={"invoice_id": fixture["invoice_id"]}
        )

    resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=admin_auth_headers,
        params={"dimension": "failed_payments", "format": "csv"},
    )
    assert resp.status_code == 200
    rows = _detail_rows(
        resp.content.decode(), ["Customer", "Email", "Failed Payment Count", "Most Recent Failure"]
    )
    assert any(r[0] == fixture["customer"]["name"] for r in rows)


def test_customers_in_past_due(client, auth_headers, admin_auth_headers, active_subscription_with_invoice):
    fixture = active_subscription_with_invoice
    with patch("app.services.payment_gateway_service.random.random", return_value=0.999):
        client.post(
            "/api/v1/payments/process", headers=auth_headers, json={"invoice_id": fixture["invoice_id"]}
        )

    resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=admin_auth_headers,
        params={"dimension": "past_due", "format": "csv"},
    )
    assert resp.status_code == 200
    rows = _detail_rows(resp.content.decode(), ["Customer", "Email", "Plan"])
    assert any(r[0] == fixture["customer"]["name"] for r in rows)


def test_customer_insights_xlsx_and_pdf_formats(client, admin_auth_headers):
    from openpyxl import load_workbook

    xlsx_resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=admin_auth_headers,
        params={"dimension": "top_paying", "format": "xlsx"},
    )
    assert xlsx_resp.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_resp.content))
    assert "Detail" in workbook.sheetnames

    pdf_resp = client.get(
        "/api/v1/reports/customer-insights",
        headers=admin_auth_headers,
        params={"dimension": "top_paying", "format": "pdf"},
    )
    assert pdf_resp.status_code == 200
    assert pdf_resp.content[:4] == b"%PDF"

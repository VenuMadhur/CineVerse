"""
Phase 5: tests for the Reports & Export module -- 5 report types x 3
formats (csv/xlsx/pdf), authorization, date-range filtering, and
invalid format handling.
"""
import csv
import io
import uuid
from unittest.mock import patch

from openpyxl import load_workbook

REPORT_PATHS = [
    "/api/v1/reports/revenue",
    "/api/v1/reports/customers",
    "/api/v1/reports/subscriptions",
    "/api/v1/reports/payments",
    "/api/v1/reports/refunds",
]


def _seed_data(client, admin_auth_headers):
    """
    Independent customer with an activated subscription, an upgrade
    (generating an invoice), a successful payment, and a partial
    refund -- enough real data for every report type to have
    something non-trivial to render.
    """
    email = f"report_{uuid.uuid4().hex[:8]}@test.com"
    client.post(
        "/api/v1/auth/register",
        json={"email": email, "full_name": "Report Data", "password": "SuperSecret123", "country": "US"},
    )
    token = client.post(
        "/api/v1/auth/login", json={"email": email, "password": "SuperSecret123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    plan_a = client.post(
        "/api/v1/plans",
        headers=admin_auth_headers,
        json={"name": f"Report Plan A {uuid.uuid4().hex[:6]}", "price": 15, "billing_interval": "monthly", "trial_period_days": 0},
    ).json()
    plan_b = client.post(
        "/api/v1/plans",
        headers=admin_auth_headers,
        json={"name": f"Report Plan B {uuid.uuid4().hex[:6]}", "price": 35, "billing_interval": "monthly"},
    ).json()
    customer = client.get("/api/v1/customers/me", headers=headers).json()
    sub = client.post(
        "/api/v1/subscriptions",
        headers=headers,
        json={"customer_id": customer["id"], "plan_id": plan_a["id"]},
    ).json()
    client.post(f"/api/v1/subscriptions/{sub['id']}/activate", headers=headers)
    client.post(
        f"/api/v1/subscriptions/{sub['id']}/change-plan",
        headers=headers,
        json={"new_plan_id": plan_b["id"]},
    )
    invoice_id = client.get(
        "/api/v1/invoices", params={"customer_id": customer["id"]}, headers=headers
    ).json()["items"][0]["id"]
    with patch("app.services.payment_gateway_service.random.random", return_value=0.0):
        client.post("/api/v1/payments/process", headers=headers, json={"invoice_id": invoice_id})
    client.post(
        "/api/v1/refunds",
        headers=headers,
        json={"invoice_id": invoice_id, "reason": "Report test refund", "amount": 3.0},
    )
    return customer, sub, invoice_id


# ---------------------------------------------------------------------
# Authorization -- every report path, both directions
# ---------------------------------------------------------------------


def test_all_report_endpoints_require_admin(client, auth_headers):
    for path in REPORT_PATHS:
        resp = client.get(path, headers=auth_headers)
        assert resp.status_code == 403, f"{path} should be forbidden for a non-admin"
        assert resp.json()["error"]["code"] == "forbidden"


def test_all_report_endpoints_require_auth(client):
    for path in REPORT_PATHS:
        resp = client.get(path)
        assert resp.status_code == 401, f"{path} should require authentication"


def test_all_report_endpoints_succeed_for_admin_as_csv(client, admin_auth_headers):
    _seed_data(client, admin_auth_headers)
    for path in REPORT_PATHS:
        resp = client.get(path, params={"format": "csv"}, headers=admin_auth_headers)
        assert resp.status_code == 200, f"{path} should succeed for an admin"
        assert resp.headers["content-type"].startswith("text/csv")


# ---------------------------------------------------------------------
# Format coverage: every report x every format
# ---------------------------------------------------------------------


def test_every_report_every_format_returns_valid_content(client, admin_auth_headers):
    _seed_data(client, admin_auth_headers)
    expected_media_type = {
        "csv": "text/csv",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "pdf": "application/pdf",
    }
    for path in REPORT_PATHS:
        for fmt in ("csv", "xlsx", "pdf"):
            resp = client.get(path, params={"format": fmt}, headers=admin_auth_headers)
            assert resp.status_code == 200, f"{path}?format={fmt} failed"
            assert resp.headers["content-type"].startswith(expected_media_type[fmt])
            assert "attachment" in resp.headers["content-disposition"]
            assert len(resp.content) > 0

            if fmt == "pdf":
                assert resp.content[:4] == b"%PDF"
            elif fmt == "xlsx":
                # A valid xlsx is a real zip archive (PK magic bytes).
                assert resp.content[:2] == b"PK"
            elif fmt == "csv":
                assert len(resp.content) > 10


def test_invalid_format_returns_422(client, admin_auth_headers):
    resp = client.get(
        "/api/v1/reports/revenue", params={"format": "json"}, headers=admin_auth_headers
    )
    assert resp.status_code == 422


# ---------------------------------------------------------------------
# Content correctness per report type
# ---------------------------------------------------------------------


def test_revenue_report_csv_reflects_real_data(client, admin_auth_headers):
    _seed_data(client, admin_auth_headers)
    resp = client.get(
        "/api/v1/reports/revenue", params={"format": "csv"}, headers=admin_auth_headers
    )
    text = resp.content.decode("utf-8")
    assert "Revenue Report" in text
    assert "Total Revenue" in text
    reader = list(csv.reader(io.StringIO(text)))
    # header row for the detail table must be present somewhere
    assert ["Month", "Revenue"] in reader


def test_customer_report_xlsx_lists_customers(client, admin_auth_headers):
    customer, _sub, _invoice_id = _seed_data(client, admin_auth_headers)
    resp = client.get(
        "/api/v1/reports/customers", params={"format": "xlsx"}, headers=admin_auth_headers
    )
    workbook = load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Summary", "Detail"]
    detail = workbook["Detail"]
    emails = [row[1].value for row in detail.iter_rows(min_row=2)]
    assert customer["email"] in emails


def test_subscription_report_reflects_created_subscription(client, admin_auth_headers):
    _customer, sub, _invoice_id = _seed_data(client, admin_auth_headers)
    resp = client.get(
        "/api/v1/reports/subscriptions", params={"format": "xlsx"}, headers=admin_auth_headers
    )
    workbook = load_workbook(io.BytesIO(resp.content))
    detail = workbook["Detail"]
    ids = [row[0].value for row in detail.iter_rows(min_row=2)]
    assert sub["id"] in ids


def test_payment_report_reflects_processed_payment(client, admin_auth_headers):
    _customer, _sub, invoice_id = _seed_data(client, admin_auth_headers)
    resp = client.get(
        "/api/v1/reports/payments", params={"format": "xlsx"}, headers=admin_auth_headers
    )
    workbook = load_workbook(io.BytesIO(resp.content))
    detail = workbook["Detail"]
    invoice_ids = [row[1].value for row in detail.iter_rows(min_row=2)]
    assert invoice_id in invoice_ids


def test_refund_report_reflects_created_refund(client, admin_auth_headers):
    _customer, _sub, invoice_id = _seed_data(client, admin_auth_headers)
    resp = client.get(
        "/api/v1/reports/refunds", params={"format": "xlsx"}, headers=admin_auth_headers
    )
    workbook = load_workbook(io.BytesIO(resp.content))
    detail = workbook["Detail"]
    invoice_ids = [row[1].value for row in detail.iter_rows(min_row=2)]
    assert invoice_id in invoice_ids


def test_revenue_report_date_range_excludes_out_of_range(client, admin_auth_headers):
    _seed_data(client, admin_auth_headers)
    resp = client.get(
        "/api/v1/reports/revenue",
        params={"format": "csv", "date_from": "2000-01-01", "date_to": "2000-01-31"},
        headers=admin_auth_headers,
    )
    text = resp.content.decode("utf-8")
    assert "Total Revenue,0" in text


def test_payment_report_date_range_excludes_out_of_range(client, admin_auth_headers):
    _customer, _sub, invoice_id = _seed_data(client, admin_auth_headers)
    resp = client.get(
        "/api/v1/reports/payments",
        params={"format": "xlsx", "date_from": "2000-01-01", "date_to": "2000-01-31"},
        headers=admin_auth_headers,
    )
    assert resp.status_code == 200
    workbook = load_workbook(io.BytesIO(resp.content))
    detail = workbook["Detail"]
    invoice_ids = [row[1].value for row in detail.iter_rows(min_row=2)]
    # The seeded payment (processed "now") must be excluded by a date
    # range entirely in the year 2000.
    assert invoice_id not in invoice_ids


def test_default_format_is_csv(client, admin_auth_headers):
    resp = client.get("/api/v1/reports/customers", headers=admin_auth_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
